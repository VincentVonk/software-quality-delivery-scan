"""Excel and PDF export functions."""

from __future__ import annotations

from io import BytesIO

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def to_excel(
    results: pd.DataFrame,
    summary: pd.DataFrame,
    improvements: pd.DataFrame,
    roadmap: pd.DataFrame,
    checklist: pd.DataFrame,
    overall_score: float,
) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame({"Metric": ["Weighted Overall Maturity Score"], "Value": [round(overall_score, 2)]}).to_excel(
            writer, sheet_name="Executive Summary", index=False
        )
        summary.to_excel(writer, sheet_name="Domain Summary", index=False)
        improvements.to_excel(writer, sheet_name="Top Improvements", index=False)
        roadmap.to_excel(writer, sheet_name="Roadmap", index=False)
        checklist.to_excel(writer, sheet_name="Evidence Checklist", index=False)
        results.to_excel(writer, sheet_name="Assessment Detail", index=False)

        workbook = writer.book
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for column_cells in worksheet.columns:
                max_length = max(len(str(cell.value or "")) for cell in column_cells)
                worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_length + 2, 12), 60)

    return buffer.getvalue()


def to_pdf(summary: pd.DataFrame, improvements: pd.DataFrame, roadmap: pd.DataFrame, overall_score: float) -> bytes:
    buffer = BytesIO()
    with PdfPages(buffer) as pdf:
        fig = plt.figure(figsize=(11.69, 8.27))
        fig.suptitle("Software Quality & Delivery Capability Scan", fontsize=20, fontweight="bold")
        fig.text(0.08, 0.82, f"Weighted overall maturity score: {overall_score:.2f} / 5", fontsize=16)
        fig.text(0.08, 0.76, "Traffic-light interpretation: Red < 2.8, Amber 2.8-3.99, Green >= 4.0", fontsize=11)
        ax = fig.add_axes([0.08, 0.16, 0.86, 0.52])
        colors = summary["Status_Color"].tolist()
        ax.barh(summary["Domain"], summary["Average_Score"], color=colors)
        ax.set_xlim(0, 5)
        ax.set_xlabel("Average score")
        ax.grid(axis="x", alpha=0.25)
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        fig, ax = plt.subplots(figsize=(11.69, 8.27))
        ax.axis("off")
        ax.set_title("Top 5 Improvement Areas", fontsize=16, fontweight="bold", pad=18)
        table_data = improvements[["Domain", "Question ID", "Score", "Question"]].copy()
        table_data["Question"] = table_data["Question"].str.wrap(70)
        table = ax.table(cellText=table_data.values, colLabels=table_data.columns, loc="center", cellLoc="left")
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.scale(1, 1.8)
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

        fig, ax = plt.subplots(figsize=(11.69, 8.27))
        ax.axis("off")
        ax.set_title("Roadmap", fontsize=16, fontweight="bold", pad=18)
        roadmap_table = roadmap[["Horizon", "Domain", "Current Score", "Recommended Action"]].head(8).copy()
        roadmap_table["Recommended Action"] = roadmap_table["Recommended Action"].str.wrap(55)
        table = ax.table(cellText=roadmap_table.values, colLabels=roadmap_table.columns, loc="center", cellLoc="left")
        table.auto_set_font_size(False)
        table.set_fontsize(8)
        table.scale(1, 1.9)
        pdf.savefig(fig, bbox_inches="tight")
        plt.close(fig)

    return buffer.getvalue()
